"""Write a synthetic «فهرست کارکنان با پست قبلی» workbook (same layout as the real one, fake people).

Matches the national IDs of Template/fake personnel workbook.xlsx.
Run:  .venv/bin/python tools/make_sample_posts.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
HEADERS = ["ردیف", "شماره ملی", "نام", "نام خانوادگی", "نوع استخدام", "وضعیت اشتغال",
           "عنوان پست سازمانی در ابلاغ", "واحد سازمانی در ابلاغ", "تاریخ آخرین انتصاب",
           "عنوان پست سازمانی قبلی\nدر ابلاغ"]
ROWS = [
    ("۱", "۰۱۵۰۰۹۹۰۰۰", "حنانه", "محمدی", "پیمانی", "شاغل", "منشی", "اداره امور هیات های تخصصی", "۱۴۰۵/۰۸/۱۲",
     "ماشین نویس (دادنامه)\nمتصدی امور دفتری و بایگانی (۱۳۹۶/۰۸/۲۹)"),
    ("۲", "۰۱۵۰۰۹۹۰۰۱", "مریم", "احمدی", "رسمی", "شاغل", "کارشناس امور قضایی", "دفتر شعبه سی ونهم", "۱۴۰۵/۰۶/۳۱",
     "متصدی امور دفتری و بایگانی (۱۳۸۸/۰۴/۲۴)\nکارشناس امور حقوقی (۱۴۰۰/۱۱/۰۴)"),
    ("۳", "۰۱۵۰۰۹۹۰۰۲", "فاطمه", "ابراهیمی", "رسمی آزمایشی", "شاغل", "کارشناس امور حقوقی", "دفتر شعبه چهل و پنجم", "۱۴۰۴/۱۱/۲۵", "-"),
    ("۴", "۰۱۵۰۰۹۹۰۰۳", "سما", "اخلاقی", "پیمانی", "مرخصی استحقاقی - دفاتر", "متصدی امور دفتری", "اداره اسناد و بایگانی", "۱۴۰۵/۰۶/۰۷",
     "متصدی امور دفتری و بایگانی (۱۳۹۷/۱۰/۱۹)"),
    ("۵", "۰۱۵۰۰۹۹۰۰۴", "نیلوفر", "ترابی", "رسمی", "شاغل", "مدیر دفتر", "اداره امور هیات های تخصصی", "۱۴۰۵/۰۸/۱۷",
     "ماشین نویس\nمدیر اجرای احکام (۱۴۰۰/۱۱/۰۴)"),
]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_view.rightToLeft = True
    ws.append([])                                       # row 1 empty, like the real file
    ws.append(HEADERS)
    for r in ROWS:
        ws.append(list(r))
    head = PatternFill("solid", start_color="B4C6E7", end_color="B4C6E7")
    for c in ws[2]:
        c.font, c.fill = Font(name="B Titr", size=10.5, bold=True), head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.font = Font(name="B Nazanin", size=11)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, w in zip("ABCDEFGHIJ", (6, 14, 10, 14, 12, 14, 22, 28, 14, 42)):
        ws.column_dimensions[col].width = w
    for out in (ROOT / "Template" / "fake previous posts workbook.xlsx",
                ROOT / "personnel_evaluation_report" / "resources" / "sample_previous_posts.xlsx"):
        wb.save(out)
    print("written")


if __name__ == "__main__":
    main()
