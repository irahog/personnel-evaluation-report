"""Loading and validating the source personnel workbook into memory.

The workbook is read once into an immutable in-memory snapshot. Nothing is written
to disk and the source file is never modified.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import os
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from . import jalali
from .normalize import clean_text, header_key, normalize_national_id
from .paths import SOURCE_MAPPING


class SourceError(Exception):
    """Error with a Persian, user-facing message that contains no personal data."""


class NeedsSheetChoice(SourceError):
    def __init__(self, candidates: list[str]):
        super().__init__("بیش از یک برگه با ستون‌های لازم پیدا شد. لطفاً برگهٔ درست را انتخاب کنید.")
        self.candidates = candidates


@dataclass(frozen=True)
class FieldSpec:
    id: str
    headers: tuple[str, ...]
    type: str
    required: bool


@dataclass(frozen=True)
class Mapping:
    fields: tuple[FieldSpec, ...]
    header_scan_rows: int

    @property
    def required(self) -> list[FieldSpec]:
        return [f for f in self.fields if f.required]


def load_mapping(path: Path = SOURCE_MAPPING) -> Mapping:
    data = json.loads(path.read_text(encoding="utf-8"))
    fields = tuple(
        FieldSpec(f["id"], tuple(f["headers"]), f["type"], bool(f["required"]))
        for f in data["fields"]
    )
    ids = [f.id for f in fields]
    if len(ids) != len(set(ids)):
        raise ValueError("duplicate field id in source mapping")
    return Mapping(fields, int(data.get("header_scan_rows", 15)))


@dataclass
class Employee:
    excel_row: int
    values: dict[str, str]                      # field id → display text ('' = empty in source)
    issues: list[str] = field(default_factory=list)  # Persian notes, no personal values

    @property
    def key(self) -> str:
        return normalize_national_id(self.values.get("national_id", ""))

    @property
    def full_name(self) -> str:
        return " ".join(p for p in (self.values.get("first_name", ""), self.values.get("last_name", "")) if p)


@dataclass
class LoadedWorkbook:
    path: str
    file_name: str
    size: int
    mtime_ns: int
    sha256: str
    loaded_at: dt.datetime
    sheet_name: str
    header_row: int
    columns: dict[str, str]                      # field id → column letter
    employees: list[Employee]
    by_key: dict[str, list[Employee]]
    warnings: list[str]                          # load-level, Persian, no personal values

    @property
    def snapshot_id(self) -> str:
        return self.sha256[:16]

    @property
    def duplicate_count(self) -> int:
        return sum(1 for v in self.by_key.values() if len(v) > 1)

    def is_stale(self) -> bool:
        """True when the file on disk differs from the loaded snapshot (or is gone)."""
        try:
            st = os.stat(self.path)
        except OSError:
            return True
        return st.st_mtime_ns != self.mtime_ns or st.st_size != self.size

    def find(self, national_id_input: str) -> list[Employee]:
        return list(self.by_key.get(normalize_national_id(national_id_input), []))


# ── reading ──────────────────────────────────────────────────────────────

def read_file_bytes(path: str) -> tuple[bytes, os.stat_result]:
    p = Path(path)
    if not p.exists():
        raise SourceError("فایل انتخاب‌شده پیدا نشد. ممکن است جابه‌جا یا حذف شده باشد.")
    if p.suffix.lower() == ".xls":
        raise SourceError("قالب قدیمی ‎.xls‎ پشتیبانی نمی‌شود. فایل را در اکسل با قالب ‎.xlsx‎ ذخیره کنید.")
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        raise SourceError("فایل انتخاب‌شده فایل اکسل (‎.xlsx‎) نیست.")
    try:
        st = p.stat()
        data = p.read_bytes()
    except PermissionError:
        raise SourceError("دسترسی به فایل ممکن نیست. اگر فایل در برنامهٔ دیگری قفل شده، آن را ببندید و دوباره تلاش کنید.")
    except OSError:
        raise SourceError("خواندن فایل با خطا مواجه شد.")
    return data, st


def _open(data: bytes, *, data_only: bool):
    try:
        return load_workbook(io.BytesIO(data), data_only=data_only)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError):
        raise SourceError("فایل اکسل خراب است یا قابل خواندن نیست.")


def _find_header(ws, mapping: Mapping) -> tuple[int, dict[str, int], list[str]] | None:
    """Find the first row (within the scan window) containing every required header.

    Returns (row, {field_id: column_index}, problems) or None.
    """
    wanted = {f.id: {header_key(h) for h in f.headers} for f in mapping.fields}
    for r in range(1, min(ws.max_row, mapping.header_scan_rows) + 1):
        keys = {}
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                keys.setdefault(header_key(v), []).append(c)
        found: dict[str, int] = {}
        problems: list[str] = []
        for f in mapping.fields:
            cols = sorted({c for k in wanted[f.id] for c in keys.get(k, [])})
            if len(cols) == 1:
                found[f.id] = cols[0]
            elif len(cols) > 1:
                problems.append(f"ستون «{f.headers[0]}» بیش از یک بار تکرار شده است.")
        if all(f.id in found for f in mapping.required):
            return r, found, problems
    return None


def _missing_required(ws, mapping: Mapping) -> list[str]:
    """For error messages: required headers missing from the best-matching row."""
    best: list[str] | None = None
    for r in range(1, min(ws.max_row, mapping.header_scan_rows) + 1):
        keys = {header_key(v) for v in (ws.cell(r, c).value for c in range(1, ws.max_column + 1))
                if isinstance(v, str)}
        missing = [f.headers[0] for f in mapping.required
                   if not any(header_key(h) in keys for h in f.headers)]
        if best is None or len(missing) < len(best):
            best = missing
    return best or [f.headers[0] for f in mapping.required]


def _cell_text(value, ftype: str, issues: list[str], label: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return clean_text(value)
    if isinstance(value, (dt.datetime, dt.date)):
        if ftype == "date":
            d = value.date() if isinstance(value, dt.datetime) else value
            try:
                converted = jalali.from_gregorian(d).format()
            except ValueError:
                issues.append(f"مقدار «{label}» تاریخ میلادی خارج از محدوده است.")
                return ""
            issues.append(f"«{label}» در اکسل به صورت تاریخ میلادی ذخیره شده بود و به شمسی تبدیل شد.")
            return converted
        return clean_text(value)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (int, float)):
        text = str(value)
        if ftype == "national_id":
            if isinstance(value, int) and len(text) < 10:
                issues.append("شماره ملی در اکسل به صورت عدد ذخیره شده و صفرهای ابتدایی آن دوباره افزوده شد.")
                text = text.zfill(10)
        elif ftype == "phone" and isinstance(value, int):
            issues.append("تلفن همراه در اکسل به صورت عدد ذخیره شده و ممکن است صفر ابتدایی آن حذف شده باشد.")
        return text
    if ftype == "multiline":           # keep the author's line breaks (one previous post per line)
        lines = [clean_text(line) for line in str(value).replace("\r", "").split("\n")]
        return "\n".join(line for line in lines if line)
    text = clean_text(value)
    if ftype == "date" and text and text != "-" and jalali.parse(text) is None:
        issues.append(f"مقدار «{label}» تاریخ شمسی معتبر (سال/ماه/روز) نیست.")
    return text


def load_source(path: str, sheet_name: str | None = None, mapping: Mapping | None = None) -> LoadedWorkbook:
    mapping = mapping or load_mapping()
    data, st = read_file_bytes(path)
    wb = _open(data, data_only=True)

    if sheet_name and sheet_name not in wb.sheetnames:
        raise SourceError("برگهٔ انتخاب‌شده در این فایل وجود ندارد.")
    sheets = [wb[sheet_name]] if sheet_name else wb.worksheets
    matches = []
    for ws in sheets:
        hit = _find_header(ws, mapping)
        if hit:
            matches.append((ws, hit))
    if not matches:
        if len(wb.worksheets) == 1 or sheet_name:
            ws = sheets[0]
            missing = "، ".join(f"«{h}»" for h in _missing_required(ws, mapping))
            raise SourceError(f"ستون‌های لازم در فایل پیدا نشد: {missing}")
        raise SourceError("در هیچ‌یک از برگه‌های فایل، ستون‌های لازم (مانند «شماره ملی»، «نام» و «نام خانوادگی») پیدا نشد.")
    if len(matches) > 1:
        raise NeedsSheetChoice([ws.title for ws, _ in matches])

    ws, (header_row, cols, problems) = matches[0]
    if problems:
        raise SourceError(" ".join(problems))

    warnings: list[str] = []
    if ws.sheet_state != "visible":
        warnings.append("برگهٔ اطلاعات پرسنلی در فایل پنهان است.")
    hidden_cols = [f.headers[0] for f in mapping.fields if f.id in cols
                   and ws.column_dimensions[get_column_letter(cols[f.id])].hidden]
    if hidden_cols:
        warnings.append("این ستون‌ها در فایل پنهان هستند ولی خوانده شدند: " + "، ".join(hidden_cols))
    missing_optional = [f.headers[0] for f in mapping.fields if not f.required and f.id not in cols]
    if missing_optional:
        warnings.append("این ستون‌های اختیاری پیدا نشدند: " + "، ".join(missing_optional))

    # Formula cells without a cached value would read as empty — detect and report.
    formula_ws = _open(data, data_only=False)[ws.title]
    labels = {f.id: f.headers[0] for f in mapping.fields}
    types = {f.id: f.type for f in mapping.fields}

    employees: list[Employee] = []
    empty_id_rows: list[int] = []
    uncached = 0
    for r in range(header_row + 1, ws.max_row + 1):
        raw = {fid: ws.cell(r, c).value for fid, c in cols.items()}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in raw.values()):
            continue
        issues: list[str] = []
        values = {}
        for fid, c in cols.items():
            v = raw[fid]
            if v is None:
                fv = formula_ws.cell(r, c).value
                if isinstance(fv, str) and fv.startswith("="):
                    uncached += 1
                    issues.append(f"«{labels[fid]}» فرمول است و مقدار محاسبه‌شده در فایل ذخیره نشده است.")
            values[fid] = _cell_text(v, types[fid], issues, labels[fid])
        emp = Employee(r, values, issues)
        if not emp.key:
            empty_id_rows.append(r)
            continue
        employees.append(emp)

    by_key: dict[str, list[Employee]] = {}
    for e in employees:
        by_key.setdefault(e.key, []).append(e)

    if empty_id_rows:
        warnings.append(f"{len(empty_id_rows)} ردیف بدون شماره ملی نادیده گرفته شد (ردیف‌های اکسل: "
                        + "، ".join(map(str, empty_id_rows[:20])) + ("…" if len(empty_id_rows) > 20 else "") + ").")
    dup = [k for k, v in by_key.items() if len(v) > 1]
    if dup:
        warnings.append(f"{len(dup)} شماره ملی تکراری در فایل وجود دارد. برای این افراد گزارش ساخته نمی‌شود تا فایل اصلاح شود.")
    bad_len = sum(1 for k in by_key if not (k.isdigit() and len(k) == 10))
    if bad_len:
        warnings.append(f"{bad_len} شماره ملی در فایل ۱۰ رقمی نیست.")
    if uncached:
        warnings.append(f"{uncached} سلول فرمول‌دار بدون مقدار ذخیره‌شده است. فایل را یک بار در اکسل باز و ذخیره کنید.")
    if not employees:
        raise SourceError("هیچ ردیف کارمندی در برگهٔ انتخاب‌شده پیدا نشد.")

    return LoadedWorkbook(
        path=str(Path(path).resolve()),
        file_name=Path(path).name,
        size=st.st_size,
        mtime_ns=st.st_mtime_ns,
        sha256=hashlib.sha256(data).hexdigest(),
        loaded_at=dt.datetime.now(),
        sheet_name=ws.title,
        header_row=header_row,
        columns={fid: get_column_letter(c) for fid, c in cols.items()},
        employees=employees,
        by_key=by_key,
        warnings=warnings,
    )
