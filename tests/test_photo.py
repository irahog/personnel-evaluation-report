"""Employee photo: processing, upload, and placement in every output (synthetic images only)."""
from __future__ import annotations

import base64
import io
import json

import pytest
from openpyxl import load_workbook
from PIL import Image

from personnel_evaluation_report import export_excel, export_md
from personnel_evaluation_report.paths import TEMPLATE_MAP
from personnel_evaluation_report.photo import PhotoError, fit, process_upload
from personnel_evaluation_report.report import apply_form, new_report
from personnel_evaluation_report.source import load_source
from tests.conftest import person


def jpeg(w=300, h=400, color="steelblue", exif_orientation=None) -> bytes:
    img = Image.new("RGB", (w, h), color)
    img.paste("red", (0, 0, w // 3, h // 3))       # marker in the top-left corner
    out = io.BytesIO()
    kw = {}
    if exif_orientation:
        exif = Image.Exif()
        exif[0x0112] = exif_orientation
        exif[0x8825] = {2: (35.0, 41.0, 0.0)}      # fake GPS block — must be stripped
        kw["exif"] = exif
    img.save(out, "JPEG", **kw)
    return out.getvalue()


def test_process_keeps_aspect_and_strips_metadata():
    p = process_upload(jpeg(3000, 4000, exif_orientation=1))
    assert (p.width, p.height) == (900, 1200)
    assert Image.open(io.BytesIO(p.jpeg)).getexif() == {}


def test_exif_rotation_is_applied():
    p = process_upload(jpeg(400, 300, exif_orientation=6))   # camera held upright: rotate 90°
    assert (p.width, p.height) == (300, 400)


def test_png_transparency_becomes_white():
    img = Image.new("RGBA", (200, 200), (0, 0, 0, 0))
    out = io.BytesIO()
    img.save(out, "PNG")
    p = process_upload(out.getvalue())
    assert Image.open(io.BytesIO(p.jpeg)).getpixel((100, 100))[0] > 240


# ASCII ids: test ids are printed to the console (Windows consoles may not handle Persian), and the
# 16 MB payload is built inside the test so it never becomes part of a test id.
@pytest.mark.parametrize("case,msg", [("empty", "انتخاب نشده"), ("not-image", "معتبر"), ("too-large", "۱۵ مگابایت")],
                         ids=["empty", "not-image", "too-large"])
def test_bad_uploads(case, msg):
    data = {"empty": b"", "not-image": b"not an image", "too-large": b"x" * (16 * 1024 * 1024)}[case]
    with pytest.raises(PhotoError, match=msg):
        process_upload(data)


def test_too_small():
    with pytest.raises(PhotoError, match="کوچک"):
        process_upload(jpeg(40, 50))


def test_fit_contain_never_distorts():
    x, y, w, h = fit(200, 500, 300, 400, pad=0)
    assert abs(w / h - 300 / 400) < 1e-9 and w == 200 and y > 0
    x, y, w, h = fit(500, 200, 300, 400)
    assert abs(w / h - 0.75) < 1e-9 and h == 200 and x > 0


@pytest.fixture
def report_with_photo(make_wb):
    wb = load_source(str(make_wb([person(1)])))
    r = new_report(wb, wb.employees[0])
    apply_form(r, {})
    r.photo = process_upload(jpeg(300, 400))
    return r


def test_excel_photo_inside_box_without_distortion(report_with_photo):
    ws = load_workbook(io.BytesIO(export_excel.build_single_workbook(report_with_photo, "B Nazanin"))).active
    assert len(ws._images) == 1
    img = ws._images[0]
    ext = img.anchor.ext
    assert abs(ext.width / ext.height - 0.75) < 0.01            # aspect ratio preserved
    m = json.loads(TEMPLATE_MAP.read_text(encoding="utf-8"))
    assert m["images"]["photo"] == "Z6:AB19"                     # own section, far left of the sheet
    frm = img.anchor._from
    assert 25 <= frm.col <= 27 and 5 <= frm.row <= 18            # anchored inside the box
    assert ws[m["fields"]["photo"]].value is None                # placeholder removed


def test_excel_without_photo_keeps_placeholder(make_wb):
    wb = load_source(str(make_wb([person(1)])))
    r = new_report(wb, wb.employees[0])
    apply_form(r, {})
    ws = load_workbook(io.BytesIO(export_excel.build_single_workbook(r, "B Nazanin"))).active
    assert not ws._images and ws["Z6"].value == "محل عکس" and ws["Z4"].value == "عکس کارمند"


def test_shared_workbook_keeps_earlier_photos(report_with_photo, make_wb, tmp_path):
    shared = tmp_path / "shared.xlsx"
    export_excel.append_to_shared(str(shared), report_with_photo, "B Nazanin")
    wb = load_source(str(make_wb([person(2)], name="b.xlsx")))
    r2 = new_report(wb, wb.employees[0])
    apply_form(r2, {})
    r2.photo = process_upload(jpeg(500, 500, "green"))
    export_excel.append_to_shared(str(shared), r2, "B Nazanin")       # must not be refused
    out = load_workbook(shared)
    assert [len(s._images) for s in out.worksheets] == [1, 1]


def test_markdown_embeds_photo(report_with_photo):
    md = export_md.render(report_with_photo)
    assert "### عکس کارمند" in md
    b64 = base64.b64encode(report_with_photo.photo.jpeg).decode()
    assert f"![عکس کارمند](data:image/jpeg;base64,{b64})" in md


def test_upload_flow_and_isolation(client, make_wb):
    client.load(make_wb())
    client.post("/search", {"national_id": "0000000001"})
    r = client.post("/report/edit", {"photo": (io.BytesIO(jpeg()), "x.jpg")}, content_type="multipart/form-data")
    assert r.status_code == 302
    img = client.get("/report/photo")
    assert img.mimetype == "image/jpeg" and img.headers["Cache-Control"] == "no-store"
    view = client.get("/report/view").get_data(as_text=True)
    assert 'id="photo-img"' in view and "/report/photo" in view
    # A different employee starts without the previous photo.
    client.post("/search", {"national_id": "0000000002"})
    assert client.get("/report/photo").status_code == 404


def test_upload_error_and_remove(client, make_wb):
    client.load(make_wb())
    client.post("/search", {"national_id": "0000000001"})
    r = client.post("/report/edit", {"photo": (io.BytesIO(b"nope"), "x.jpg")}, content_type="multipart/form-data")
    assert r.status_code == 422 and "عکس معتبری نیست" in r.get_data(as_text=True)
    client.post("/report/edit", {"photo": (io.BytesIO(jpeg()), "x.jpg")}, content_type="multipart/form-data")
    assert client.get("/report/photo").status_code == 200
    client.post("/report/edit", {"photo_remove": "1"})
    assert client.get("/report/photo").status_code == 404
