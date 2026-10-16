"""Normalization, Jalali dates, workbook loading, lookup and mapping."""
from __future__ import annotations

import datetime as dt

import pytest
from openpyxl import load_workbook

from personnel_evaluation_report import jalali
from personnel_evaluation_report.normalize import header_key, normalize_national_id, to_ascii_digits, to_fa_digits
from personnel_evaluation_report.source import NeedsSheetChoice, SourceError, load_mapping, load_source
from tests.conftest import HEADERS, person, write_workbook


# ── normalization ─────────────────────────────────────────────────────────
def test_digits_persian_arabic_ascii():
    assert to_ascii_digits("۰۱۲۳۴۵۶۷۸۹") == "0123456789"
    assert to_ascii_digits("٠١٢٣٤٥٦٧٨٩") == "0123456789"
    assert to_fa_digits("0012") == "۰۰۱۲"


@pytest.mark.parametrize("raw", ["۰۰۱۲۳۴۵۶۷۸", "٠٠١٢٣٤٥٦٧٨", "0012345678", " 001-234-5678 ", "‏0012345678‎"])
def test_national_id_normalization_keeps_leading_zeros(raw):
    assert normalize_national_id(raw) == "0012345678"


def test_header_key_explicit_rules():
    assert header_key("محل خدمت ـ اداره کل") == header_key("محل خدمت - اداره کل") == header_key("محل خدمت – اداره کل")
    assert header_key(" وضعيت  تاهل ") == header_key("وضعیت تاهل")       # Arabic yeh, extra spaces
    assert header_key("نام") != header_key("نام خانوادگی")                  # no partial matching


# ── Jalali ────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("greg,jal", [
    (dt.date(1921, 3, 21), (1300, 1, 1)),
    (dt.date(2024, 3, 20), (1403, 1, 1)),
    (dt.date(2025, 3, 20), (1403, 12, 30)),   # 1403 is a leap year
    (dt.date(2025, 3, 21), (1404, 1, 1)),
    (dt.date(2026, 10, 16), (1405, 7, 24)),
    (dt.date(1979, 2, 11), (1357, 11, 22)),
])
def test_gregorian_to_jalali(greg, jal):
    d = jalali.from_gregorian(greg)
    assert (d.year, d.month, d.day) == jal


def test_jalali_parse_and_validate():
    assert jalali.parse("۱۴۰۵/۰۷/۲۴") == jalali.JalaliDate(1405, 7, 24)
    assert jalali.parse("1403/12/30") is not None          # leap
    assert jalali.parse("1404/12/30") is None              # not leap
    assert jalali.parse("1404/07/31") is None              # months 7–11 have 30 days
    assert jalali.parse("1390/13/01") is None
    assert jalali.parse("0150099003") is None
    assert jalali.parse("1361-1-9") == jalali.JalaliDate(1361, 1, 9)


# ── mapping ───────────────────────────────────────────────────────────────
def test_mapping_preserves_exact_source_headers():
    m = load_mapping()
    first = {f.headers[0] for f in m.fields}
    assert set(HEADERS) <= first, "every header from the brief must appear verbatim in the mapping"
    assert {f.id for f in m.required} >= {"national_id", "first_name", "last_name", "work_dept", "divan_entry_date"}


# ── loading ───────────────────────────────────────────────────────────────
def test_load_basic(make_wb):
    wb = load_source(str(make_wb()))
    assert wb.sheet_name == "Sheet1" and wb.header_row == 2 and len(wb.employees) == 5
    assert wb.columns["national_id"] == "G"
    assert wb.duplicate_count == 0


def test_header_in_row_one_and_reordered_extra_columns(tmp_path):
    hdrs = list(reversed(HEADERS)) + ["ستون اضافه"]
    p = write_workbook(tmp_path / "x.xlsx", [person(1, **{"ستون اضافه": "x"})], headers=hdrs, title_rows=0)
    wb = load_source(str(p))
    assert wb.header_row == 1
    e = wb.find("0000000001")[0]
    assert e.values["first_name"] == "نام1" and e.values["work_dept"] == "بخش آزمایشی 1"


def test_header_variants_from_fake_workbook(tmp_path):
    variants = {"محل خدمت ـ اداره کل": "محل خدمت - اداره کل", "محل خدمت ـ اداره/بخش": "محل خدمت - اداره/بخش",
                "روزسابقه خدمت دولتی": "روز سابقه خدمت دولتی", "ماه سابقه خدمت دولتی": "ماه سابقه خدمتی دولتی",
                "سال سابقه خدمت دولتی": "سال سابقه خدمتی دولتی"}
    hdrs = [variants.get(h, h) for h in HEADERS]
    rows = [{variants.get(k, k): v for k, v in person(1).items()}]
    wb = load_source(str(write_workbook(tmp_path / "x.xlsx", rows, headers=hdrs)))
    assert wb.find("0000000001")[0].values["work_dept"] == "بخش آزمایشی 1"
    assert not any("اختیاری" in w for w in wb.warnings)


def test_missing_required_column_is_reported(tmp_path):
    hdrs = [h for h in HEADERS if h != "شماره ملی"]
    p = write_workbook(tmp_path / "x.xlsx", [person(1)], headers=hdrs)
    with pytest.raises(SourceError, match="شماره ملی"):
        load_source(str(p))


def test_missing_optional_column_is_a_warning(tmp_path):
    hdrs = [h for h in HEADERS if h != "وضعیت اشتغال"]
    wb = load_source(str(write_workbook(tmp_path / "x.xlsx", [person(1)], headers=hdrs)))
    assert any("وضعیت اشتغال" in w for w in wb.warnings)


def test_sheet_detection_skips_other_sheets(tmp_path):
    p = write_workbook(tmp_path / "x.xlsx", [person(1)], sheet="کارکنان",
                       extra_sheets={"یادداشت": (["توضیح"], [{"توضیح": "متن"}])})
    wb_ = load_workbook(p)
    wb_.move_sheet("یادداشت", offset=-1)          # make the wrong sheet first
    wb_.save(p)
    assert load_source(str(p)).sheet_name == "کارکنان"


def test_ambiguous_sheets_require_choice(tmp_path):
    p = write_workbook(tmp_path / "x.xlsx", [person(1)], sheet="الف",
                       extra_sheets={"ب": (HEADERS, [person(2)])})
    with pytest.raises(NeedsSheetChoice) as e:
        load_source(str(p))
    assert set(e.value.candidates) == {"الف", "ب"}
    assert load_source(str(p), "ب").find("0000000002")


def test_hidden_sheet_is_flagged(tmp_path):
    p = write_workbook(tmp_path / "x.xlsx", [person(1)], extra_sheets={"دیگر": (["x"], [])})
    wb_ = load_workbook(p)
    wb_["Sheet1"].sheet_state = "hidden"
    wb_.active = 1
    wb_.save(p)
    assert any("پنهان" in w for w in load_source(str(p)).warnings)


def test_duplicate_national_ids_detected(make_wb):
    rows = [person(1), person(2), person(3, **{"شماره ملی": "0000000001"})]
    wb = load_source(str(make_wb(rows)))
    assert wb.duplicate_count == 1
    assert len(wb.find("۰۰۰۰۰۰۰۰۰۱")) == 2
    assert any("تکراری" in w for w in wb.warnings)


def test_numeric_national_id_gets_zero_padding_with_note(make_wb):
    rows = [person(1, **{"شماره ملی": 12345678, "تلفن همراه": 9120000000})]
    wb = load_source(str(make_wb(rows)))
    e = wb.find("0012345678")[0]
    assert e.values["national_id"] == "0012345678"
    assert any("صفرهای ابتدایی" in i for i in e.issues)
    assert e.values["mobile"] == "9120000000"                       # not fabricated
    assert any("تلفن همراه" in i for i in e.issues)


def test_exact_lookup_only(make_wb):
    wb = load_source(str(make_wb()))
    assert wb.find("000000000") == []          # prefix is not a match
    assert wb.find("00000000011") == []
    assert len(wb.find("٠٠٠٠٠٠٠٠٠٣")) == 1


def test_empty_cells_and_rows(make_wb):
    rows = [person(1, **{"رشته تحصیلی": None}), {}, person(2, **{"شماره ملی": None})]
    wb = load_source(str(make_wb(rows)))
    assert len(wb.employees) == 1
    assert wb.find("0000000001")[0].values["field_of_study"] == ""
    assert any("بدون شماره ملی" in w for w in wb.warnings)


def test_excel_date_cell_is_converted_and_flagged(make_wb):
    rows = [person(1, **{"تاریخ تولد": dt.datetime(1986, 5, 5)})]
    e = load_source(str(make_wb(rows))).find("0000000001")[0]
    assert e.values["birth_date"] == "۱۳۶۵/۰۲/۱۵"
    assert any("میلادی" in i for i in e.issues)


def test_invalid_text_date_is_flagged_not_changed(make_wb):
    rows = [person(1, **{"تاریخ تولد": "۰۱۵۰۰۹۹۰۰۳"})]
    e = load_source(str(make_wb(rows))).find("0000000001")[0]
    assert e.values["birth_date"] == "۰۱۵۰۰۹۹۰۰۳"
    assert any("تاریخ شمسی معتبر" in i for i in e.issues)


def test_formula_cells(tmp_path, make_wb):
    p = make_wb()
    wb_ = load_workbook(p)
    ws = wb_.active
    ws["M3"] = '="رشته"&"ی فرمولی"'      # formula without cached value (openpyxl never computes)
    wb_.save(p)
    w = load_source(str(p))
    assert w.find("0000000001")[0].values["field_of_study"] == ""
    assert any("فرمول" in x for x in w.warnings)


def test_bad_files(tmp_path):
    (tmp_path / "a.xlsx").write_bytes(b"not a zip")
    with pytest.raises(SourceError, match="خراب"):
        load_source(str(tmp_path / "a.xlsx"))
    (tmp_path / "b.xls").write_bytes(b"x")
    with pytest.raises(SourceError, match="xls"):
        load_source(str(tmp_path / "b.xls"))
    with pytest.raises(SourceError, match="پیدا نشد"):
        load_source(str(tmp_path / "missing.xlsx"))
    (tmp_path / "c.csv").write_text("x")
    with pytest.raises(SourceError):
        load_source(str(tmp_path / "c.csv"))


def test_paths_with_spaces_and_persian(tmp_path):
    d = tmp_path / "پوشه با فاصله"
    d.mkdir()
    p = write_workbook(d / "فایل پرسنلی ۱۴۰۵.xlsx", [person(1)])
    assert load_source(str(p)).file_name == "فایل پرسنلی ۱۴۰۵.xlsx"


def test_source_file_is_not_modified(make_wb):
    p = make_wb()
    before = p.read_bytes()
    load_source(str(p))
    assert p.read_bytes() == before


def test_staleness_detection(make_wb):
    p = make_wb()
    wb = load_source(str(p))
    assert not wb.is_stale()
    write_workbook(p, [person(9)])
    import os
    st = os.stat(p)
    os.utime(p, ns=(st.st_atime_ns, wb.mtime_ns + 10_000_000))
    assert wb.is_stale()


def test_bundled_fake_workbook_loads():
    from personnel_evaluation_report.paths import SAMPLE_WORKBOOK
    wb = load_source(str(SAMPLE_WORKBOOK))
    assert len(wb.employees) == 5 and wb.header_row == 1
    assert wb.find("۰۱۵۰۰۹۹۰۰۰")[0].values["national_id"].startswith("۰")
