"""Per-report table colours and text styles: parsing, CSS, Excel, isolation."""
from __future__ import annotations

import io
import json

from openpyxl import load_workbook

from personnel_evaluation_report import export_excel, styles
from personnel_evaluation_report.paths import TEMPLATE_MAP
from personnel_evaluation_report.report import SECTIONS, apply_form, new_report
from personnel_evaluation_report.source import load_source
from tests.conftest import person

F = styles.field_name


def form_for(key, part, **kw):
    d = {F(key, part, "present"): "1", F(key, part, "fg"): kw.get("fg", styles.DEFAULTS[part].fg),
         F(key, part, "bg"): kw.get("bg", styles.DEFAULTS[part].bg), F(key, part, "size"): kw.get("size", "normal")}
    for flag in ("bold", "italic", "underline"):
        if kw.get(flag, getattr(styles.DEFAULTS[part], flag)):
            d[F(key, part, flag)] = "1"
    return d


def test_every_table_section_has_a_key():
    titled = [s.title for s in SECTIONS if s.kind not in ("banner", "photo")]
    assert set(titled) == set(styles.SECTION_KEYS)
    regions = json.loads(TEMPLATE_MAP.read_text(encoding="utf-8"))["sections"]
    assert set(regions) == set(styles.SECTION_KEYS.values())
    assert all(regions[k]["header"] and regions[k]["body"] for k in regions)


def test_defaults_store_nothing():
    form = {}
    for key in styles.SECTION_KEYS.values():
        for part in styles.PARTS:
            form.update(form_for(key, part))
    st, errors = styles.parse_form(form)
    assert st == {} and errors == []
    assert styles.parse_form({}) == ({}, [])          # controls absent → defaults


def test_custom_values_and_validation():
    st, errors = styles.parse_form({**form_for("discipline", "header", fg="#AA0000", bg="#ffee00", italic=True),
                                    **form_for("history", "body", bold=True, size="large"),
                                    **form_for("satisfaction", "body", fg="red")})
    assert st["discipline"]["header"] == styles.PartStyle("#aa0000", "#ffee00", True, True, False, "normal")
    assert st["history"]["body"].bold and st["history"]["body"].scale == 1.2
    assert "satisfaction" not in st and len(errors) == 1 and "رضایت سنجی" in errors[0]


def test_css_targets_section():
    st, _ = styles.parse_form(form_for("competency", "body", fg="#123456", underline=True, size="small"))
    css = styles.css(st)
    assert 'section.rsec[data-sec="competency"] tbody td' in css and "color:#123456" in css
    assert "text-decoration:underline" in css and "font-size:0.85em" in css
    assert "</style" not in css


def make_report(make_wb):
    wb = load_source(str(make_wb([person(1)])))
    r = new_report(wb, wb.employees[0])
    apply_form(r, {})
    return r


def test_excel_applies_styles_but_keeps_fixed_colours(make_wb):
    r = make_report(make_wb)
    r.styles, _ = styles.parse_form({**form_for("discipline", "header", fg="#aa0000", bg="#ffee00", italic=True),
                                     **form_for("discipline", "body", bg="#e0f0ff", size="large"),
                                     **form_for("employment", "body", bg="#00ff00")})
    ws = load_workbook(io.BytesIO(export_excel.build_single_workbook(r, "B Mitra"))).active
    title = ws["K23"]
    assert title.font.color.rgb.endswith("AA0000") and title.font.italic and title.fill.fgColor.rgb.endswith("FFEE00")
    assert title.font.name == "B Mitra" and title.font.size == 13
    body = ws["L25"]                                   # «ورود وخروج به موقع»
    assert body.fill.fgColor.rgb.endswith("E0F0FF") and body.font.size == 14.4
    assert ws["O31"].fill.fgColor.rgb.endswith("FFC000")     # «تذکر» keeps its official colour
    assert ws["S7"].fill.fgColor.rgb.endswith("FFFF00")      # yellow box untouched by employment body bg
    assert ws["N6"].fill.fgColor.rgb.endswith("00FF00")
    assert ws["B25"].fill.fgColor.rgb in ("00000000", None) or not ws["B25"].fill.fgColor.rgb.endswith("E0F0FF")


def test_web_flow_and_reset_on_new_employee(client, make_wb):
    client.load(make_wb())
    client.post("/search", {"national_id": "0000000001"})
    page = client.get("/report/edit").get_data(as_text=True)
    assert page.count('class="style-panel') == 9 and page.count("data-sec=") == 9
    client.post("/report/edit", form_for("history", "header", bg="#ff0000"))
    css = client.get("/report/style.css")
    assert css.mimetype == "text/css" and "#ff0000" in css.get_data(as_text=True)
    view = client.get("/report/view").get_data(as_text=True)
    assert "/report/style.css" in view and 'data-sec="history"' in view
    assert "style-panel" not in view                    # controls only in the form
    client.post("/search", {"national_id": "0000000002"})
    assert client.get("/report/style.css").get_data(as_text=True).strip() == ""


def test_style_css_requires_report(client):
    assert client.get("/report/style.css").status_code == 302
