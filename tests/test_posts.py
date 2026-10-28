"""Optional second workbook «فهرست کارکنان با پست قبلی»: loading, matching, prefill, routes."""
from __future__ import annotations

import io
import json
import os

from openpyxl import load_workbook

from personnel_evaluation_report import export_excel
from personnel_evaluation_report.paths import POSTS_MAPPING, TEMPLATE_MAP
from personnel_evaluation_report.report import apply_form, new_report, prefill_from_posts
from personnel_evaluation_report.source import load_mapping, load_source
from tests.conftest import person, write_workbook

POST_HEADERS = ["ردیف", "شماره ملی", "نام", "نام خانوادگی", "نوع استخدام", "وضعیت اشتغال",
                "عنوان پست سازمانی در ابلاغ", "واحد سازمانی در ابلاغ", "تاریخ آخرین انتصاب",
                "عنوان پست سازمانی قبلی\nدر ابلاغ"]


def post(i, prev="متصدی امور دفتری (۱۳۹۶/۰۸/۲۹)\nکارشناس (۱۴۰۰/۱۱/۰۴)", date="۱۴۰۰/۰۴/۲۰", nid=None):
    from personnel_evaluation_report.normalize import to_fa_digits
    return {"ردیف": str(i), "شماره ملی": nid or to_fa_digits(f"0000000{i:03d}"), "نام": f"نام{i}",
            "نام خانوادگی": f"خانوادگی{i}", "تاریخ آخرین انتصاب": date, "عنوان پست سازمانی قبلی\nدر ابلاغ": prev}


def posts_file(tmp_path, rows, name="posts.xlsx"):
    return write_workbook(tmp_path / name, rows, headers=POST_HEADERS)


def load_posts(path):
    return load_source(str(path), mapping=load_mapping(POSTS_MAPPING))


def report_for(make_wb, i=1):
    wb = load_source(str(make_wb([person(1), person(2), person(3)])))
    return new_report(wb, wb.find(f"000000000{i}")[0])


def test_loader_handles_wrapped_header_and_multiline_values(tmp_path):
    pw = load_posts(posts_file(tmp_path, [post(1)]))
    assert pw.header_row == 2 and pw.columns["previous_posts"] == "J" and pw.columns["last_appointment"] == "I"
    assert pw.find("0000000001")[0].values["previous_posts"] == "متصدی امور دفتری (۱۳۹۶/۰۸/۲۹)\nکارشناس (۱۴۰۰/۱۱/۰۴)"


def test_prefill_by_national_id(make_wb, tmp_path):
    r = report_for(make_wb, 2)
    prefill_from_posts(r, load_posts(posts_file(tmp_path, [post(1), post(2, prev="ماشین نویس", date="۱۳۹۷/۰۵/۰۱")])))
    assert r.manual["previous_posts"] == "ماشین نویس" and r.manual["post_suspension"] == "۱۳۹۷/۰۵/۰۱"
    assert set(r.prefilled) == {"previous_posts", "post_suspension"} and r.posts_notes == []


def test_not_found_and_duplicates_are_reported_not_guessed(make_wb, tmp_path):
    r = report_for(make_wb, 3)
    prefill_from_posts(r, load_posts(posts_file(tmp_path, [post(1)])))
    assert not r.manual.get("previous_posts") and "پیدا نشد" in r.posts_notes[0]
    r = report_for(make_wb, 1)
    prefill_from_posts(r, load_posts(posts_file(tmp_path, [post(1), post(1, prev="دیگر")], name="d.xlsx")))
    assert not r.manual.get("previous_posts") and "تکراری" in r.posts_notes[0]
    assert any("تکراری" in w for w in r.warnings())


def test_prefill_never_overwrites_user_input_and_stays_editable(make_wb, tmp_path):
    r = report_for(make_wb, 1)
    r.manual["previous_posts"] = "نوشتهٔ کاربر"
    prefill_from_posts(r, load_posts(posts_file(tmp_path, [post(1)])))
    assert r.manual["previous_posts"] == "نوشتهٔ کاربر" and r.manual["post_suspension"] == "۱۴۰۰/۰۴/۲۰"
    apply_form(r, {"previous_posts": "اصلاح‌شده", "post_suspension": "۱۴۰۱/۰۱/۰۱"})
    assert r.manual["previous_posts"] == "اصلاح‌شده" and r.manual["post_suspension"] == "۱۴۰۱/۰۱/۰۱"


def test_missing_required_posts_column(tmp_path):
    import pytest
    from personnel_evaluation_report.source import SourceError
    hdrs = [h for h in POST_HEADERS if h != "تاریخ آخرین انتصاب"]
    with pytest.raises(SourceError, match="تاریخ آخرین انتصاب"):
        load_posts(write_workbook(tmp_path / "x.xlsx", [post(1)], headers=hdrs))


def test_excel_keeps_line_breaks(make_wb, tmp_path):
    r = report_for(make_wb, 1)
    prefill_from_posts(r, load_posts(posts_file(tmp_path, [post(1)])))
    apply_form(r, {"previous_posts": r.manual["previous_posts"], "post_suspension": r.manual["post_suspension"]})
    ws = load_workbook(io.BytesIO(export_excel.build_single_workbook(r, "B Nazanin"))).active
    cells = json.loads(TEMPLATE_MAP.read_text(encoding="utf-8"))["fields"]
    assert ws[cells["previous_posts"]].value == "متصدی امور دفتری (۱۳۹۶/۰۸/۲۹)\nکارشناس (۱۴۰۰/۱۱/۰۴)"
    assert ws[cells["post_suspension"]].value == "۱۴۰۰/۰۴/۲۰"


def test_web_flow(client, make_wb, tmp_path):
    client.load(make_wb())
    client.answers.append(str(posts_file(tmp_path, [post(1), post(2)])))
    body = client.post("/posts/pick", follow_redirects=True).get_data(as_text=True)
    assert "فایل پست‌های قبلی" in body and "بارگذاری شد" in body
    client.post("/search", {"national_id": "0000000001"})
    page = client.get("/report/edit").get_data(as_text=True)
    assert "کارشناس (۱۴۰۰/۱۱/۰۴)" in page and "از فایل پست‌ها" in page and "۱۴۰۰/۰۴/۲۰" in page
    client.post("/report/edit", {"previous_posts": "ویرایش کاربر", "post_suspension": "۱۴۰۰/۰۴/۲۰"})
    assert "ویرایش کاربر" in client.get("/report/view").get_data(as_text=True)


def test_posts_loaded_after_report_fills_empty_fields(client, make_wb, tmp_path):
    client.load(make_wb())
    client.post("/search", {"national_id": "0000000002"})
    client.answers.append(str(posts_file(tmp_path, [post(2, prev="پست قبلی دو")])))
    client.post("/posts/pick")
    assert "پست قبلی دو" in client.get("/report/edit").get_data(as_text=True)


def test_optional_and_clear(client, make_wb, tmp_path):
    client.load(make_wb())
    client.post("/search", {"national_id": "0000000001"})
    assert "از فایل پست‌ها" not in client.get("/report/edit").get_data(as_text=True)   # works without the file
    client.answers.append(str(posts_file(tmp_path, [post(1)])))
    client.post("/posts/pick")
    assert "کنار گذاشته شد" in client.post("/posts/clear", follow_redirects=True).get_data(as_text=True)


def test_rejects_personnel_file_as_posts_file(client, make_wb):
    p = make_wb()
    client.load(p)
    client.answers.append(str(p))
    assert "همان فایل اطلاعات پرسنلی" in client.post("/posts/pick", follow_redirects=True).get_data(as_text=True)


def test_stale_posts_file_blocks_search(client, make_wb, tmp_path):
    client.load(make_wb())
    p = posts_file(tmp_path, [post(1)])
    client.answers.append(str(p))
    client.post("/posts/pick")
    st = os.stat(p)
    os.utime(p, ns=(st.st_atime_ns, st.st_mtime_ns + 5_000_000_000))
    body = client.post("/search", {"national_id": "0000000001"}, follow_redirects=True).get_data(as_text=True)
    assert "فایل پست‌های قبلی پس از بارگذاری تغییر کرده" in body
    assert client.get("/report/edit").status_code == 302


def test_bundled_sample_posts_matches_sample_personnel():
    from personnel_evaluation_report.paths import RESOURCES, SAMPLE_WORKBOOK
    pw = load_posts(RESOURCES / "resources" / "sample_previous_posts.xlsx")
    wb = load_source(str(SAMPLE_WORKBOOK))
    assert {e.key for e in pw.employees} == {e.key for e in wb.employees}
