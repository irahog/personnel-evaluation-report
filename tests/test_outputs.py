"""Report model, validation, and consistency of Excel / Markdown / HTML outputs."""
from __future__ import annotations

import io
import json
import zipfile

import pytest
from openpyxl import load_workbook

from personnel_evaluation_report import export_excel, export_md, fonts
from personnel_evaluation_report.paths import TEMPLATE_MAP
from personnel_evaluation_report.report import FIELDS, SECTIONS, TableSection, all_cells, apply_form, new_report
from personnel_evaluation_report.source import load_source
from tests.conftest import person

MANUAL = {
    "children_daughters": "۲", "children_sons": "1", "service_history": "سابقهٔ آزمایشی\nخط دوم",
    "comp_1_score": "۱۸", "comp_2_score": "17/5", "comp_total": "۳۵٫۵", "disc_7_pledge": "تعهد آزمایشی",
    "recommend_3_mark": "1", "hist_2_rank": "اول", "hist_average": "90", "sat_2_desc": "توضیح | با خط عمودی",
    "final_score": "۸۸٫۵", "evaluator_signature": "ارزیاب آزمایشی", "evaluation_date": "۱۴۰۵/۰۷/۲۴",
}


@pytest.fixture
def report(make_wb):
    wb = load_source(str(make_wb([person(1, **{"وضعیت تاهل": "متاهل"}), person(2)])))
    r = new_report(wb, wb.find("0000000001")[0])
    assert apply_form(r, MANUAL)
    return r


def test_spec_matches_template_map():
    cells = json.loads(TEMPLATE_MAP.read_text(encoding="utf-8"))["fields"]
    for fid, f in FIELDS.items():
        if f.input == "marital":
            assert {"marital_married_mark", "marital_single_mark"} <= cells.keys()
        else:
            assert fid in cells, fid
    used = {c for s in SECTIONS if isinstance(s, TableSection) for c in all_cells(s)}
    used |= {f for s in SECTIONS if hasattr(s, "fields") for f in s.fields}
    assert used == set(FIELDS), "every field appears exactly in the section layout"


def test_auto_mapping(report):
    a = report.auto
    assert a["full_name"] == "نام1 خانوادگی1"
    assert a["work_location"] == "بخش آزمایشی 1"          # محل خدمت ← اداره/بخش
    assert a["start_date"] == "۱۳۹۵/۰۷/۰۱"                # ← تاریخ ورود به دیوان
    assert a["mobile"] == "۰۹۱۲۰۰۰۰۰۰۰" and a["national_id"] == "۰۰۰۰۰۰۰۰۰۱"
    assert report.marital == "married" and report.marital_from_source


def test_no_calculation(report):
    # Totals are whatever the user typed — never computed.
    assert report.number("comp_total") == 35.5
    assert report.number("comp_1_score") + report.number("comp_2_score") == 35.5
    r2 = report
    apply_form(r2, {**MANUAL, "comp_total": ""})
    assert r2.manual["comp_total"] == "" and r2.number("comp_total") is None


def test_validation_errors(report):
    ok = apply_form(report, {**MANUAL, "comp_1_score": "هجده", "children_sons": "-1",
                             "evaluation_date": "1404/12/30"})
    assert not ok
    assert set(report.errors) == {"comp_1_score", "children_sons", "evaluation_date"}
    assert not report.validated


def test_unknown_marital_status_needs_manual_choice(make_wb):
    wb = load_source(str(make_wb([person(1, **{"وضعیت تاهل": "معیل"})])))
    r = new_report(wb, wb.employees[0])
    assert r.marital == "" and not r.marital_from_source
    assert any("معیل" in w for w in r.warnings())
    apply_form(r, {"marital": "single"})
    assert r.marital == "single"


def test_marital_from_source_cannot_be_overridden(report):
    apply_form(report, {**MANUAL, "marital": "single"})
    assert report.marital == "married"


def test_markdown(report):
    md = export_md.render(report)
    assert md.startswith("# کارنامه پرسنلی")
    order = [s.title for s in SECTIONS]
    positions = [md.index(t) for t in order]
    assert positions == sorted(positions), "section order preserved"
    assert "| کدملی | ۰۰۰۰۰۰۰۰۰۱ |" in md
    assert "☑ متأهل" in md and "☐ مجرد" in md
    assert "توضیح \\| با خط عمودی" in md                    # table-safe escaping
    assert "سابقهٔ آزمایشی<br>خط دوم" in md
    assert "۱۷/۵" in md and "۸۸/۵" in md
    md.encode("utf-8")


def test_excel_single_workbook(report, tmp_path):
    data = export_excel.build_single_workbook(report, "B Mitra")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["نام1 خانوادگی1"]
    ws = wb.active
    cells = json.loads(TEMPLATE_MAP.read_text(encoding="utf-8"))["fields"]
    assert ws[cells["national_id"]].value == "۰۰۰۰۰۰۰۰۰۱"
    assert ws[cells["work_location"]].value == "بخش آزمایشی 1"
    assert ws[cells["marital_married_mark"]].value == "☑" and ws[cells["marital_single_mark"]].value == "☐"
    assert ws[cells["comp_2_score"]].value == 17.5 and ws[cells["final_score"]].value == 88.5
    assert ws[cells["children_daughters"]].value == 2
    assert ws[cells["recommend_3_mark"]].value == "☑" and ws[cells["recommend_1_mark"]].value == "☐"
    assert ws[cells["comp_3_score"]].value is None             # empty stays empty
    assert ws.sheet_view.rightToLeft
    assert ws.page_setup.orientation == "landscape" and ws.print_area
    assert ws.merged_cells.ranges and ws.row_breaks.brk
    names = {c.font.name for row in ws.iter_rows() for c in row if c.value is not None}
    assert names <= {"B Mitra", fonts.SYMBOL_FONT}
    assert ws[cells["full_name"]].font.name == "B Mitra"
    # only the current employee's report — never the whole personnel list
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "نام2" not in text


def test_outputs_contain_same_values(report, client, make_wb):
    """Every manual and auto value appears identically in Markdown, HTML and Excel."""
    data = export_excel.build_single_workbook(report, "B Nazanin")
    ws = load_workbook(io.BytesIO(data)).active
    cells = json.loads(TEMPLATE_MAP.read_text(encoding="utf-8"))["fields"]
    md = export_md.render(report)
    from personnel_evaluation_report.webapp import create_app
    app = create_app()
    with app.test_request_context("/"):
        from flask import g, render_template
        from personnel_evaluation_report.webapp import SessionState
        g.s = SessionState(report=report)
        html = render_template("view.html", r=report)
    for fid, f in FIELDS.items():
        if f.input in ("marital", "check"):
            continue
        t = report.text(fid)
        if not t:
            continue
        line = t.split("\n")[0]
        assert export_md._md(line) in md, fid
        assert line.replace("|", "&#124;") in html or line in html, fid
        xv = ws[cells[fid]].value
        if f.input in ("number", "count"):
            assert xv == report.number(fid), fid
        else:
            assert xv == t, fid


def test_sheet_title_sanitizing():
    assert export_excel.sheet_title("علی/رضا: [آزمایش]*?") == "علی رضا آزمایش"
    assert len(export_excel.sheet_title("الف" * 50)) == 31
    assert export_excel.sheet_title("  ") == "کارمند"


def test_shared_workbook_append_replace_and_new_name(report, make_wb, tmp_path):
    shared = tmp_path / "مشترک.xlsx"
    res = export_excel.append_to_shared(str(shared), report, "B Nazanin")
    assert res.sheet == "نام1 خانوادگی1" and not res.replaced
    with pytest.raises(export_excel.SameEmployeeExists):
        export_excel.append_to_shared(str(shared), report, "B Nazanin")
    res = export_excel.append_to_shared(str(shared), report, "B Nazanin", replace=True)
    assert res.replaced
    # Another employee with the same name → new sheet with suffix, nothing overwritten.
    wb = load_source(str(make_wb([person(5, **{"نام": "نام1", "نام خانوادگی": "خانوادگی1"})], name="o.xlsx")))
    other = new_report(wb, wb.employees[0])
    apply_form(other, {})
    res = export_excel.append_to_shared(str(shared), other, "B Nazanin")
    assert res.sheet == "نام1 خانوادگی1 (2)"
    assert load_workbook(shared).sheetnames == ["نام1 خانوادگی1", "نام1 خانوادگی1 (2)"]
    assert not list(tmp_path.glob("~per-report-*")), "temporary files are cleaned up"


def test_shared_workbook_with_drawings_is_refused(report, tmp_path):
    p = tmp_path / "old.xlsx"
    from openpyxl import Workbook
    Workbook().save(p)
    with zipfile.ZipFile(p, "a") as z:
        z.writestr("xl/drawings/drawing1.xml", "<x/>")
    before = p.read_bytes()
    with pytest.raises(export_excel.ExportError, match="شکل"):
        export_excel.append_to_shared(str(p), report, "B Nazanin")
    assert p.read_bytes() == before


def test_shared_workbook_corrupt_and_wrong_type(report, tmp_path):
    p = tmp_path / "bad.xlsx"
    p.write_bytes(b"garbage")
    with pytest.raises(export_excel.ExportError):
        export_excel.append_to_shared(str(p), report, "B Nazanin")
    with pytest.raises(export_excel.ExportError):
        export_excel.append_to_shared(str(tmp_path / "x.xls"), report, "B Nazanin")


def test_font_decision(monkeypatch):
    monkeypatch.setattr(fonts, "is_installed", lambda f: f == "Times New Roman")
    d = export_excel.decide_font("B Nazanin", "Times New Roman")
    assert d.used == "Times New Roman" and "B Nazanin" in d.warning
    monkeypatch.setattr(fonts, "is_installed", lambda f: False)
    d = export_excel.decide_font("B Mitra", "Times New Roman")
    assert d.used == "B Mitra" and d.warning           # never silently swapped
    monkeypatch.setattr(fonts, "is_installed", lambda f: True)
    assert export_excel.decide_font("B Mitra", "Times New Roman").warning == ""


def test_bundled_fonts_present_and_cover_persian():
    files = fonts.bundled_files()
    assert len(files) == 6 and all(p.exists() for p in files.values())
    css = (fonts.FONTS_DIR.parent / "fonts.css").read_text(encoding="utf-8")
    for name in files:
        assert f"fonts/{name}" in css


def test_install_user_fonts(tmp_path, monkeypatch):
    registered = []
    monkeypatch.setattr(fonts, "user_font_dir", lambda: tmp_path / "fonts")
    monkeypatch.setattr(fonts, "_register_windows_font", lambda name, path: registered.append(name))
    monkeypatch.setattr(fonts, "_broadcast_font_change", lambda: None)
    done = fonts.install_user_fonts()
    assert sorted(done) == sorted(fonts.bundled_files())
    assert (tmp_path / "fonts" / "BNazanin.ttf").read_bytes() == fonts.bundled_files()["BNazanin.ttf"].read_bytes()
    assert fonts.install_user_fonts() == []                 # idempotent
    assert set(fonts.FULL_NAMES) == set(fonts.bundled_files())


def test_install_fonts_route(client, tmp_path, monkeypatch):
    monkeypatch.setattr(fonts, "user_font_dir", lambda: tmp_path / "f")
    monkeypatch.setattr(fonts, "_register_windows_font", lambda name, path: None)
    monkeypatch.setattr(fonts, "_broadcast_font_change", lambda: None)
    r = client.post("/settings/install-fonts", follow_redirects=True)
    assert "نصب شد" in r.get_data(as_text=True)
    r = client.post("/settings/install-fonts", follow_redirects=True)
    assert "از قبل نصب" in r.get_data(as_text=True)
